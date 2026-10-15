import io
import csv
from decimal import Decimal
from django.utils import timezone
from django.http import HttpResponse
from django.db.models import Sum, Count, Avg
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib.pagesizes import letter, A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

from .models import Campaign, Student, Place, ReceiptBook, Trip, TripStudent, Expense, CommissionCalculation, CommissionPayout
from .services import FinancialService, InventoryService


class ReportGenerator:
    @staticmethod
    def get_campaign_data(campaign: Campaign) -> dict:
        fin = FinancialService.get_campaign_financial_summary(campaign)
        inv = InventoryService.get_campaign_inventory_summary(campaign)
        return {
            'financial': fin,
            'inventory': inv,
            'campaign': campaign,
        }

    @staticmethod
    def export_excel(campaign: Campaign, report_type: str = 'full') -> HttpResponse:
        wb = openpyxl.Workbook()
        # Remove default sheet
        wb.remove(wb.active)

        # Styles
        header_fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        title_font = Font(name="Calibri", size=14, bold=True, color="111827")
        bold_font = Font(name="Calibri", size=10, bold=True)
        regular_font = Font(name="Calibri", size=10)
        total_fill = PatternFill(start_color="F3F4F6", end_color="F3F4F6", fill_type="solid")
        thin_border = Border(
            left=Side(style='thin', color='E5E7EB'),
            right=Side(style='thin', color='E5E7EB'),
            top=Side(style='thin', color='E5E7EB'),
            bottom=Side(style='thin', color='E5E7EB')
        )

        # Sheet 1: Campaign Summary
        ws_sum = wb.create_sheet(title="Campaign Summary")
        ws_sum.append(["ZAINUSSUNNA ACADEMY - CALENDAR FUND MANAGEMENT SYSTEM"])
        ws_sum.append([f"Campaign: {campaign.name} ({campaign.year})", f"Status: {campaign.status.upper()}"])
        ws_sum.append([f"Generated: {timezone.now().strftime('%d %b %Y, %I:%M %p')}"])
        ws_sum.append([])

        fin = FinancialService.get_campaign_financial_summary(campaign)
        inv = InventoryService.get_campaign_inventory_summary(campaign)

        summary_rows = [
            ["Metric", "Value", "Notes"],
            ["Total Calendars Printed", inv['printed_quantity'], "Initial opening batch"],
            ["Calendars Sold", inv['calendars_sold'], "Confirmed settled sales"],
            ["Calendars Currently in Field", inv['calendars_in_field'], "On active trips"],
            ["Calendars Returned Unsold", inv['calendars_returned'], "Safely returned to inventory"],
            ["Available Stock in Hand", inv['available_stock'], "Printed - Sold - In Field + Adjustments"],
            ["Total Collections", float(fin['total_collections']), "Cash + Google Pay"],
            ["  - Cash Collections", float(fin['cash_collections']), ""],
            ["  - Google Pay / UPI", float(fin['gpay_collections']), ""],
            ["Total Trip Expenses", float(fin['total_trip_expenses']), "Bus fare + Food expenses"],
            ["  - Bus Fare Advanced", float(fin['bus_fare']), ""],
            ["  - Food Expenses", float(fin['food_expense']), ""],
            ["General Campaign Expenses", float(fin['general_expenses']), "Excluding trip deductions"],
            ["  - Printing Cost", float(fin['printing_expense']), f"₹{campaign.cost_price_per_calendar}/calendar"],
            ["Total Commission Earned", float(fin['total_commission_earned']), f"{campaign.commission_percentage}% on student net collection"],
            ["Total Commission Paid", float(fin['total_commission_paid']), "Actual disbursements recorded"],
            ["Outstanding Commission Payable", float(fin['outstanding_commission']), ""],
            ["Net Campaign Profit", float(fin['net_profit']), "Collections - Trip Expenses - General Expenses - Commission Earned"],
            ["Completed / Settled Trips", fin['settled_trips_count'], ""],
            ["Active Trips in Progress", fin['open_trips_count'], ""],
            ["Trips Requiring Review", fin['review_trips_count'], ""],
        ]

        for idx, row in enumerate(summary_rows):
            ws_sum.append(row)
            row_idx = ws_sum.max_row
            if idx == 0:
                for col_idx in range(1, 4):
                    cell = ws_sum.cell(row=row_idx, column=col_idx)
                    cell.fill = header_fill
                    cell.font = header_font
            elif "Profit" in str(row[0]) or "Total Collections" in str(row[0]):
                for col_idx in range(1, 4):
                    cell = ws_sum.cell(row=row_idx, column=col_idx)
                    cell.fill = total_fill
                    cell.font = bold_font

        # Sheet 2: Trips
        ws_trips = wb.create_sheet(title="Trips")
        trips_headers = [
            "Trip #", "Date", "Place", "Students", "Status", "Review",
            "Issued", "Sold", "Returned", "Cash (₹)", "GPay (₹)", "Total Coll. (₹)",
            "Bus Fare (₹)", "Food Exp. (₹)", "Receipt Book", "Leaf Range"
        ]
        ws_trips.append(trips_headers)
        for col_idx in range(1, len(trips_headers) + 1):
            cell = ws_trips.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font

        trips = Trip.objects.filter(campaign=campaign).select_related('place', 'receipt_book').prefetch_related('trip_students__student')
        for t in trips:
            students_str = ", ".join(s.student.name for s in t.trip_students.all())
            leaf_str = f"{t.receipt_starting_leaf or ''} - {t.receipt_ending_leaf or ''}"
            ws_trips.append([
                t.id, str(t.trip_date), t.place.name, students_str, t.status, t.review_status,
                t.calendars_issued, t.calendars_sold, t.calendars_returned,
                float(t.cash_collection), float(t.gpay_collection), float(t.total_collection),
                float(t.bus_fare), float(t.food_expense),
                t.receipt_book.label if t.receipt_book else "", leaf_str
            ])

        # Sheet 3: Students & Commission
        ws_students = wb.create_sheet(title="Students & Commission")
        students_headers = [
            "Student Name", "Trips", "Gross Share (₹)", "Bus Share (₹)", "Food Share (₹)",
            "Net Collection (₹)", "Comm. Rate (%)", "Earned Comm. (₹)", "Paid (₹)", "Outstanding (₹)"
        ]
        ws_students.append(students_headers)
        for col_idx in range(1, len(students_headers) + 1):
            cell = ws_students.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font

        students = Student.objects.all()
        for s in students:
            calc = FinancialService.recalculate_student_commission(campaign, s)
            trips_count = TripStudent.objects.filter(trip__campaign=campaign, trip__status='settled', student=s).count()
            if trips_count > 0 or calc.gross_collection > 0:
                paid = CommissionPayout.objects.filter(campaign=campaign, student=s).aggregate(t=Sum('amount'))['t'] or Decimal('0.00')
                rem = max(Decimal('0.00'), calc.calculated_commission - paid)
                ws_students.append([
                    s.name, trips_count, float(calc.gross_collection), float(calc.fare_share),
                    float(calc.food_share), float(calc.net_collection), float(calc.rate),
                    float(calc.calculated_commission), float(paid), float(rem)
                ])

        # Sheet 4: Expenses
        ws_exp = wb.create_sheet(title="Expenses")
        exp_headers = ["ID", "Category", "Amount (₹)", "Date", "Description", "Voided?"]
        ws_exp.append(exp_headers)
        for col_idx in range(1, len(exp_headers) + 1):
            cell = ws_exp.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font

        expenses = Expense.objects.filter(campaign=campaign)
        for e in expenses:
            ws_exp.append([
                e.id, e.category, float(e.amount), str(e.date), e.description, "YES" if e.is_voided else "NO"
            ])

        # Auto-adjust column widths
        for ws in wb.worksheets:
            for col in ws.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                col_letter = openpyxl.utils.get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = min(40, max(max_len + 3, 12))

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        response = HttpResponse(
            output.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        response['Content-Disposition'] = f'attachment; filename="Zainussunna_Campaign_{campaign.year}_Report.xlsx"'
        return response

    @staticmethod
    def export_csv(campaign: Campaign, report_type: str = 'trips') -> HttpResponse:
        output = io.StringIO()
        writer = csv.writer(output)

        if report_type == 'trips':
            writer.writerow([
                "Trip ID", "Date", "Place", "Students", "Status", "Review Status",
                "Calendars Issued", "Calendars Sold", "Calendars Returned",
                "Cash Collection", "GPay Collection", "Total Collection",
                "Bus Fare", "Food Expense", "Receipt Book", "Starting Leaf", "Ending Leaf"
            ])
            trips = Trip.objects.filter(campaign=campaign).select_related('place', 'receipt_book').prefetch_related('trip_students__student')
            for t in trips:
                st_names = "; ".join(s.student.name for s in t.trip_students.all())
                writer.writerow([
                    t.id, t.trip_date, t.place.name, st_names, t.status, t.review_status,
                    t.calendars_issued, t.calendars_sold, t.calendars_returned,
                    t.cash_collection, t.gpay_collection, t.total_collection,
                    t.bus_fare, t.food_expense,
                    t.receipt_book.label if t.receipt_book else "",
                    t.receipt_starting_leaf or "", t.receipt_ending_leaf or ""
                ])
            filename = f"Zainussunna_Trips_{campaign.year}.csv"

        elif report_type == 'students':
            writer.writerow([
                "Student ID", "Name", "Settled Trips", "Gross Collection Share",
                "Bus Fare Share", "Food Expense Share", "Net Collection",
                "Commission Rate (%)", "Commission Earned", "Commission Paid", "Outstanding Payable"
            ])
            students = Student.objects.all()
            for s in students:
                calc = FinancialService.recalculate_student_commission(campaign, s)
                trips_count = TripStudent.objects.filter(trip__campaign=campaign, trip__status='settled', student=s).count()
                if trips_count > 0 or calc.gross_collection > 0:
                    paid = CommissionPayout.objects.filter(campaign=campaign, student=s).aggregate(t=Sum('amount'))['t'] or Decimal('0.00')
                    rem = max(Decimal('0.00'), calc.calculated_commission - paid)
                    writer.writerow([
                        s.id, s.name, trips_count, calc.gross_collection,
                        calc.fare_share, calc.food_share, calc.net_collection,
                        calc.rate, calc.calculated_commission, paid, rem
                    ])
            filename = f"Zainussunna_Students_{campaign.year}.csv"

        else:  # expenses
            writer.writerow(["Expense ID", "Category", "Amount", "Date", "Description", "Is Voided"])
            for e in Expense.objects.filter(campaign=campaign):
                writer.writerow([e.id, e.category, e.amount, e.date, e.description, e.is_voided])
            filename = f"Zainussunna_Expenses_{campaign.year}.csv"

        response = HttpResponse(output.getvalue(), content_type="text/csv")
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response

    @staticmethod
    def export_pdf(campaign: Campaign) -> HttpResponse:
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=A4,
            leftMargin=36,
            rightMargin=36,
            topMargin=36,
            bottomMargin=36
        )

        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            'TitleStyle',
            parent=styles['Heading1'],
            fontName='Helvetica-Bold',
            fontSize=16,
            leading=20,
            textColor=colors.HexColor('#111827'),
            alignment=1,  # Center
        )
        subtitle_style = ParagraphStyle(
            'SubTitleStyle',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=10,
            leading=14,
            textColor=colors.HexColor('#4B5563'),
            alignment=1,
        )
        section_style = ParagraphStyle(
            'SectionStyle',
            parent=styles['Heading2'],
            fontName='Helvetica-Bold',
            fontSize=12,
            leading=16,
            textColor=colors.HexColor('#1F2937'),
            spaceBefore=12,
            spaceAfter=6,
        )
        cell_style = ParagraphStyle(
            'CellStyle',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=8,
            leading=10,
        )

        elements = []

        # Header
        elements.append(Paragraph("ZAINUSSUNNA ACADEMY", title_style))
        elements.append(Paragraph(f"Calendar Fund Management System — Campaign {campaign.year} Financial Statement", subtitle_style))
        elements.append(Paragraph(f"Status: {campaign.status.upper()} | Generated on {timezone.now().strftime('%d %b %Y, %I:%M %p')}", subtitle_style))
        elements.append(Spacer(1, 16))

        fin = FinancialService.get_campaign_financial_summary(campaign)
        inv = InventoryService.get_campaign_inventory_summary(campaign)

        # Summary Table
        summary_data = [
            ["Metric", "Value", "Metric", "Value"],
            ["Total Calendars Printed", str(inv['printed_quantity']), "Total Collections (₹)", f"₹{fin['total_collections']:,.2f}"],
            ["Calendars Sold", str(inv['calendars_sold']), "  - Cash Collections (₹)", f"₹{fin['cash_collections']:,.2f}"],
            ["Calendars in Field", str(inv['calendars_in_field']), "  - GPay / UPI Collections (₹)", f"₹{fin['gpay_collections']:,.2f}"],
            ["Calendars Returned Unsold", str(inv['calendars_returned']), "Total Trip Expenses (₹)", f"₹{fin['total_trip_expenses']:,.2f}"],
            ["Stock Available in Hand", str(inv['available_stock']), "General Expenses (₹)", f"₹{fin['general_expenses']:,.2f}"],
            ["Settled Trips", str(fin['settled_trips_count']), "Printing Cost (₹)", f"₹{fin['printing_expense']:,.2f}"],
            ["Active Trips", str(fin['open_trips_count']), "Total Commission Earned (₹)", f"₹{fin['total_commission_earned']:,.2f}"],
            ["Trips Needing Review", str(fin['review_trips_count']), "Commission Paid (₹)", f"₹{fin['total_commission_paid']:,.2f}"],
            ["Break-Even Progress", f"{fin['break_even_percentage']}%", "Outstanding Commission (₹)", f"₹{fin['outstanding_commission']:,.2f}"],
            ["", "", "Net Campaign Profit (₹)", f"₹{fin['net_profit']:,.2f}"],
        ]

        t_summary = Table(summary_data, colWidths=[130, 130, 130, 130])
        t_summary.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1F2937')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 8),
            ('ALIGN', (1, 0), (1, -1), 'RIGHT'),
            ('ALIGN', (3, 0), (3, -1), 'RIGHT'),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E5E7EB')),
            ('BACKGROUND', (2, -1), (3, -1), colors.HexColor('#E5E7EB')),
            ('FONTNAME', (2, -1), (3, -1), 'Helvetica-Bold'),
        ]))
        elements.append(t_summary)
        elements.append(Spacer(1, 14))

        # Recent Settled Trips Section
        elements.append(Paragraph("Settled Trips", section_style))
        trips_table_data = [["Trip", "Date", "Place", "Sold", "Ret.", "Total (₹)", "Fare (₹)", "Food (₹)", "Status"]]
        trips = Trip.objects.filter(campaign=campaign, status='settled')[:20]
        for t in trips:
            trips_table_data.append([
                f"#{t.id}", str(t.trip_date), t.place.name[:18],
                str(t.calendars_sold), str(t.calendars_returned),
                f"₹{t.total_collection:.0f}", f"₹{t.bus_fare:.0f}", f"₹{t.food_expense:.0f}",
                t.review_status.replace('_', ' ').capitalize()
            ])

        if len(trips_table_data) > 1:
            t_trips = Table(trips_table_data, colWidths=[35, 55, 120, 40, 40, 60, 50, 50, 70])
            t_trips.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#374151')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, -1), 7),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
                ('TOPPADDING', (0, 0), (-1, -1), 3),
                ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E5E7EB')),
            ]))
            elements.append(t_trips)
        else:
            elements.append(Paragraph("No settled trips recorded yet.", cell_style))

        elements.append(Spacer(1, 14))

        # Student Commission Summary Section
        elements.append(Paragraph("Student Commission Summary", section_style))
        st_table_data = [["Student Name", "Trips", "Gross Share (₹)", "Deductions (₹)", "Net Share (₹)", "Comm. (₹)", "Paid (₹)", "Outstanding (₹)"]]
        students = Student.objects.filter(
            trip_participations__trip__campaign=campaign,
            trip_participations__trip__status='settled'
        ).distinct()[:25]

        for s in students:
            calc = FinancialService.recalculate_student_commission(campaign, s)
            trips_count = TripStudent.objects.filter(trip__campaign=campaign, trip__status='settled', student=s).count()
            paid = CommissionPayout.objects.filter(campaign=campaign, student=s).aggregate(t=Sum('amount'))['t'] or Decimal('0.00')
            rem = max(Decimal('0.00'), calc.calculated_commission - paid)
            deductions = calc.fare_share + calc.food_share
            st_table_data.append([
                s.name[:20], str(trips_count), f"₹{calc.gross_collection:.2f}", f"₹{deductions:.2f}",
                f"₹{calc.net_collection:.2f}", f"₹{calc.calculated_commission:.2f}",
                f"₹{paid:.2f}", f"₹{rem:.2f}"
            ])

        if len(st_table_data) > 1:
            t_students = Table(st_table_data, colWidths=[90, 35, 65, 65, 65, 65, 65, 70])
            t_students.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#374151')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, -1), 7),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
                ('TOPPADDING', (0, 0), (-1, -1), 3),
                ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E5E7EB')),
            ]))
            elements.append(t_students)
        else:
            elements.append(Paragraph("No student commission calculated yet.", cell_style))

        doc.build(elements)
        buffer.seek(0)

        response = HttpResponse(buffer.getvalue(), content_type='application/pdf')
        response['Content-Disposition'] = f'attachment; filename="Zainussunna_Financial_Report_{campaign.year}.pdf"'
        return response
